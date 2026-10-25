# -*- coding: utf-8 -*-
# 【購入者限定】本ファイルの再配布・転載・転売は固く禁じます。｜UI診断ディレクター スキル版 prototype r6（2026-10-26｜設計 r2・判定113）
"""check_output.py｜生成後の点検（r6）｜不合格なら AI が入力を直して作り直す

使い方:
    python3 check_output.py <入力JSON> <出力フォルダ>
        入力JSON = {"site_slug":..., "record":{...}, "diagnosis":{...}, "proposals":{...}, "visual_data":{...}}
    python3 check_output.py --pptx-only <出力フォルダ>
        入力の無い PPTX（旧版など）の表示だけを調べる（試験用）。内容・構成の点検は実施しない
    どちらも末尾に --report <JSONの保存先> を付けると、欄ごとの結果を JSON で保存する

結果は3段で表示する:
    1. 内容・構成点検（r5 の点検をすべて維持＋入力の A・B・C・文字種・record）
    2. 表示の簡易推定（幅・高さ・枠の対応・例外表・フッター・切り詰め）
       ※「基準内」は実際の描画の合格を意味しない
    3. Windows PowerPoint 目視：未実施（最終の判定は Windows の PowerPoint での目視）

終了コード: 0＝内容・構成が合格かつ 表示の推定で NG 0・未検査 0／1＝不合格（直す箇所を表示）
            2＝引数誤り／3＝--pptx-only で NG 0・未検査 0（内容は未実施のため合格ではない）

r5 の点検項目（一つも外さない）:
    1. 総評（visual_data.summary）：項目があり・空でなく・17字以内・末尾が「…」「...」でない
    2. ランクの説明（rank_label）：ランクに対応する1語（S＝卓越／A＝優良／B＝標準的／C＝要改善／D＝致命的）
    3. 件数：強み・課題とも各3件（ランクに関わらず・PPTX に載る件数）。C-1・C-3・top_issues で同じ
    4. ファイル名が英数字で3点そろう／test_*.pptx が無い
    5. PPTX にハイパーリンクが無い／メイリオ・14pt 以上
    6. PPTX の中の文字が入力と一致（C-3 の総評・C-1 のランク欄・強み・課題・top_issues の各文言）
    7. C-3 の Before/After：target_area・priority・effort と PPTX の「箇所：」／flow_summary 40字以内で PPTX にある
"""
import json, os, re, sys, glob, zipfile

ITEM_COUNT = 3  # 強み・課題の件数（builder の描画上限｜r5）
RANK_SET = ('S', 'A', 'B', 'C', 'D')
RANK_LABEL = {'S': '卓越', 'A': '優良', 'B': '標準的', 'C': '要改善', 'D': '致命的'}
SUMMARY_MAX = 17
NAME_RE = {
    'C-1': re.compile(r'^01_UI_Scorecard_[A-Za-z0-9_]+\.pptx$'),
    'C-2': re.compile(r'^02_Proposals_[A-Za-z0-9_]+\.pptx$'),
    'C-3': re.compile(r'^03_Visual_Board_[A-Za-z0-9_]+\.pptx$'),
}


def slide_texts(path):
    """PPTX の各スライドの文字（段落ごと）を返す"""
    out = []
    with zipfile.ZipFile(path) as z:
        names = sorted((n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n)),
                       key=lambda n: int(re.findall(r'\d+', n)[0]))
        for n in names:
            x = z.read(n).decode('utf-8', 'ignore')
            paras = []
            for p in re.findall(r'<a:p>.*?</a:p>|<a:p .*?</a:p>', x, flags=re.S):
                t = ''.join(re.findall(r'<a:t>([^<]*)</a:t>', p))
                if t:
                    paras.append(t.replace('&amp;', '&').replace('&lt;', '<').replace('&gt;', '>'))
            out.append(paras)
    return out


def check(data, out_dir):
    ng = []
    d, v = data.get('diagnosis', {}), data.get('visual_data', {})
    # 1. 総評
    if 'summary' not in v:
        ng.append('総評（visual_data.summary）の項目が無い → 17字以内の総評を入れる')
        s = None
    else:
        s = v.get('summary')
        if not isinstance(s, str) or not s.strip():
            ng.append('総評が空 → 17字以内の総評を入れる')
        else:
            if len(s) > SUMMARY_MAX:
                ng.append(f'総評が{len(s)}字（上限{SUMMARY_MAX}字）：「{s}」→ 意味を保って{SUMMARY_MAX}字以内に言い換える')
            if s.rstrip().endswith(('…', '...', '‥')):
                ng.append(f'総評の末尾が省略記号：「{s}」→ 省略せず言い切る形に言い換える')
    # 2・3. ランクの説明と件数
    rank = str(d.get('rank', '')).strip()
    need = ITEM_COUNT if rank in RANK_SET else None
    if need is None:
        ng.append(f'ランク「{rank}」が S/A/B/C/D のどれでもない')
    else:
        for label, src in [('C-1', d), ('C-3', v)]:
            if 'rank_label' in src and src.get('rank_label') != RANK_LABEL[rank]:
                ng.append(f'{label} のランクの説明が「{src.get("rank_label")}」→ ランク{rank}は「{RANK_LABEL[rank]}」をそのまま使う')
        if d.get('rank_label') is None:
            ng.append(f'C-1 のランクの説明（rank_label）が無い → 「{RANK_LABEL[rank]}」を入れる')
        for label, arr in [('C-1 強み', d.get('strengths')), ('C-1 課題', d.get('priority_issues')),
                           ('C-3 強み', v.get('strengths')), ('C-3 課題', v.get('priority_issues')),
                           ('C-3 top_issues', v.get('top_issues'))]:
            n = len(arr or [])
            if n != need:
                ng.append(f'{label}が{n}件（ランクに関わらず{need}件｜4件目以降は改善提案・結論に回す）')
    # 4. ファイル
    files = sorted(os.path.basename(p) for p in glob.glob(os.path.join(out_dir, '*.pptx')))
    found = {}
    for key, rx in NAME_RE.items():
        hit = [f for f in files if rx.match(f)]
        if len(hit) != 1:
            ng.append(f'{key} のファイルが{len(hit)}件（1件であるべき・英数字のファイル名）')
        else:
            found[key] = os.path.join(out_dir, hit[0])
    for f in files:
        if not re.match(r'^[A-Za-z0-9_.-]+$', f):
            ng.append(f'英数字以外のファイル名：{f}')
    tests = [f for f in files if f.startswith('test_')]
    if tests:
        ng.append(f'test_*.pptx がある：{tests}（builder の読み込み方を確認。デモが動いている）')
    # 5. リンク・フォント
    for f in files:
        with zipfile.ZipFile(os.path.join(out_dir, f)) as z:
            for n in z.namelist():
                if n.endswith('.rels') and 'hyperlink' in z.read(n).decode('utf-8', 'ignore'):
                    ng.append(f'ハイパーリンクがある：{f}（{n}）')
                if re.match(r'ppt/slides/slide\d+\.xml$', n):
                    x = z.read(n).decode('utf-8', 'ignore')
                    small = [int(sz) for sz in re.findall(r'<a:rPr[^>]*\ssz="(\d+)"', x) if int(sz) < 1400]
                    if small:
                        ng.append(f'14pt 未満の文字：{f}（{n}・{min(small)/100}pt）')
                    faces = set(re.findall(r'typeface="([^"]+)"', x)) - {'+mn-lt', '+mj-lt', '+mn-ea', '+mj-ea', '+mn-cs', '+mj-cs'}
                    other = sorted(t for t in faces if t not in ('Meiryo', 'メイリオ', 'Meiryo UI'))
                    if other:
                        ng.append(f'メイリオ以外のフォント：{f}（{n}・{other}）')
    # 6. PPTX の中身と入力の照合
    if 'C-3' in found and isinstance(s, str) and s.strip():
        paras = [p for sl in slide_texts(found['C-3']) for p in sl]
        if s not in paras:
            ng.append(f'C-3 の PPTX に総評「{s}」が見当たらない → 入力を直したら PPTX も作り直す')
    if 'C-1' in found and need is not None:
        want = f'{rank}  {d.get("rank_label") or ""}'.rstrip()
        sl = slide_texts(found['C-1'])
        hit_all = all(want in paras for paras in sl[:2]) if len(sl) >= 2 else (want in sl[0] if sl else False)
        if not hit_all:
            ng.append(f'C-1 の PPTX のランク欄が入力（「{want}」）と一致しない → 入力を直したら PPTX も作り直す')
    # 6-2. 強み・課題・top_issues の各文言が PPTX の中にあるか（判定68）
    def _norm(t):
        return re.sub(r'\s+', '', t or '')
    for key, src, items in [('C-1', d, [('強み', 'strengths'), ('課題', 'priority_issues')]),
                            ('C-3', v, [('強み', 'strengths'), ('課題', 'priority_issues'), ('top_issues', 'top_issues')])]:
        if key not in found:
            continue
        whole = _norm(''.join(p for sl in slide_texts(found[key]) for p in sl))
        for jp, field in items:
            for t in (src.get(field) or []):
                if isinstance(t, str) and t.strip() and _norm(t) not in whole:
                    ng.append(f'{key} {jp}『{t}』が PPTX に見当たらない → 入力を直したら PPTX も作り直す')
    # 7. C-3 の Before/After と行動の流れのまとめ（判定78）
    hls = v.get('highlights') or []
    c3_paras = [p for sl in slide_texts(found['C-3']) for p in sl] if 'C-3' in found else []
    c3_whole = _norm(''.join(c3_paras))
    for i, h in enumerate(hls[:3], start=1):
        if not isinstance(h, dict):
            ng.append(f'C-3 Before/After #{i} の形が違う'); continue
        ta = h.get('target_area')
        if not (isinstance(ta, str) and ta.strip()):
            if h.get('area'):
                ng.append(f'C-3 Before/After #{i}：箇所が area になっている → area ではなく target_area と書く（C-2 の対応する提案と同じ値）')
            else:
                ng.append(f'C-3 Before/After #{i}：箇所（target_area）が無い → C-2 の対応する提案の target_area を入れる')
        elif 'C-3' in found and _norm('箇所：' + ta[:20]) not in c3_whole:
            ng.append(f'C-3 Before/After #{i}：PPTX に「箇所：{ta}」が見当たらない → 入力を直したら PPTX も作り直す')
        if h.get('priority') not in ('高', '中', '低'):
            ng.append(f'C-3 Before/After #{i}：優先度（priority）が「{h.get("priority")}」→ 高・中・低のどれか（C-2 と同じ値）')
        if h.get('effort') not in ('小', '中', '大'):
            ng.append(f'C-3 Before/After #{i}：工数（effort）が「{h.get("effort")}」→ 小・中・大のどれか（C-2 と同じ値）')
    fs = v.get('flow_summary')
    if not (isinstance(fs, str) and fs.strip()):
        ng.append('C-3 行動の流れのまとめ（flow_summary）が無い・空 → 40字以内の1文を入れる')
    else:
        if len(fs) > 40:
            ng.append(f'C-3 行動の流れのまとめが{len(fs)}字（上限40字）→ 言い換える')
        elif 'C-3' in found and _norm('→ ' + fs) not in c3_whole:
            ng.append(f'C-3 の PPTX に行動の流れのまとめ「{fs}」が見当たらない → 入力を直したら PPTX も作り直す')
    return files, ng




# =====================================================================
# ここから r6 の追加（設計 r2｜3891789c7f87・判定113）。r5 の点検は上の check() のまま
# =====================================================================
import unicodedata

AUTHOR = 'UI診断ディレクター'
EMU_PX = 9525.0            # 1280×720 座標（builder の px()）
TOL = 2                    # 枠の位置の許容差（px）

# ---- §2-2 固定記号の幅（固定記号の一覧を先に当てる）----
FIXED_W = {'✓': 1.0, '✕': 1.0, '⚠': 1.0, '→': 1.0, '◎': 1.0, '◤': 1.0, '●': 1.0,
           '？': 1.0, '\u2060': 0.0}


def char_class(ch):
    """§2-1 の文字種。戻り値 (幅 or None, NG理由 or None)"""
    o = ord(ch)
    if ch in '\n\r\t':
        return None, '改行・タブは使わない'
    if 0x20 <= o <= 0x7E:
        return 0.5, None
    if 0xFF61 <= o <= 0xFF9F:
        return None, '半角カナは全角に直す'
    if (0x0300 <= o <= 0x036F) or (0xFE00 <= o <= 0xFE0F) or (0xE0100 <= o <= 0xE01EF) or (0x200B <= o <= 0x200D):
        return None, '結合文字・異体字セレクタ・ゼロ幅文字は使わない'
    if (0xE000 <= o <= 0xF8FF) or o >= 0xF0000:
        return None, '私用領域の文字は使わない'
    if o >= 0x1F000:
        return None, '絵文字は使わない'
    if unicodedata.category(ch) in ('Cc', 'Cf', 'Cs', 'Co', 'Cn'):
        return None, '制御文字・書式文字は使わない'
    if ch == '\u3000':
        return 1.0, None
    eaw = unicodedata.east_asian_width(ch)
    if eaw in ('W', 'F', 'A'):
        return 1.0, None
    return None, f'未対応の文字種（U+{o:04X}）'


def disp_w(s):
    """入力文字列の表示幅（C の数え方）。(幅, NG理由のリスト)"""
    w, bad = 0.0, []
    for ch in s:
        cw, why = char_class(ch)
        if why:
            bad.append(f'「{ch}」(U+{ord(ch):04X})：{why}')
        else:
            w += cw
    return w, bad


def builder_w(s):
    """builder の _truncate_full の数え方（ord<128 を0.5・それ以外 1.0）"""
    return sum(0.5 if ord(c) < 128 else 1.0 for c in s)


def pptx_char_w(ch):
    """PPTX 上の文字の幅（固定記号を先に当て、次に文字種）。定義外は None"""
    if ch in FIXED_W:
        return FIXED_W[ch]
    cw, why = char_class(ch)
    return None if why else cw


# ---- §1-2 上限（A＝len／B＝builder 換算／C＝表示換算）----
LIM = {
    'service_name':  dict(A=30, C=11),
    'input_type':    dict(C=10),
    'comment':       dict(A=40, C=40),
    'strength':      dict(A=40, B=28),
    'issue_top':     dict(A=20, B=20, C=20),     # priority_issues・top_issues
    'conclusion':    dict(A=60, C=114),
    'p_title':       dict(A=30, B=36, C=32),
    'target_area':   dict(A=25, B=20, C=10),
    'p_issue':       dict(A=25, C=28),
    'before_after':  dict(A=40, B=38, C=24),
    'p_summary':     dict(A=120, B=120, C=114),
    'purpose':       dict(A=20, C=14),
    'diag_date':     dict(C=14),
    'v_summary':     dict(A=17, B=17, C=17),
    'sec_label':     dict(A=20, B=12, C=9),
    'sec_desc':      dict(B=16, C=10),
    'flow_label':    dict(B=6, C=6),
    'flow_note':     dict(B=14, C=9),
    'flow_summary':  dict(A=40, B=40, C=31),
    'hl_title':      dict(A=30, B=30, C=27),
    'direction':     dict(A=120, B=120, C=112),
    'top_issue':     dict(A=40, C=20),
}
LIM_NAME = {'A': 'len', 'B': 'builder換算', 'C': '表示換算'}


def check_input_r6(data):
    """入力の段階の点検（§1・§2・§5）。NG のリストを返す"""
    ng, notes = [], []
    d = data.get('diagnosis') or {}
    p = data.get('proposals') or {}
    v = data.get('visual_data') or {}

    def chk(label, s, key):
        if s is None:
            return
        if not isinstance(s, str):
            ng.append(f'{label}：文字列でない'); return
        w, bad = disp_w(s)
        for b in bad:
            ng.append(f'{label}：{b}')
        lim = LIM[key]
        if 'A' in lim and len(s) > lim['A']:
            ng.append(f'{label}：len {len(s)}（上限 {lim["A"]}）「{s}」→ 言い換える')
        if 'B' in lim and builder_w(s) > lim['B']:
            ng.append(f'{label}：builder 換算 {builder_w(s):g}（上限 {lim["B"]}）「{s}」→ 言い換える（builder が「…」で切るため）')
        if 'C' in lim and not bad and w > lim['C']:
            ng.append(f'{label}：表示上限超過 {w:g}（上限 {lim["C"]}）「{s}」→ 言い換える')
        return w

    for lbl, src in (('C-1', d), ('C-2', p), ('C-3', v)):
        chk(f'{lbl} service_name', src.get('service_name'), 'service_name')
    names = [src.get('service_name') for src in (d, p, v)]
    if len(set(names)) != 1:
        ng.append(f'service_name が3か所で一致しない：{names}')
    for lbl, src in (('C-1', d), ('C-3', v)):
        chk(f'{lbl} input_type', src.get('input_type'), 'input_type')
    for i, sc in enumerate(d.get('scores') or [], 1):
        if isinstance(sc, dict):
            chk(f'C-1 scores#{i} comment', sc.get('comment'), 'comment')
    for lbl, src in (('C-1', d), ('C-3', v)):
        for i, s in enumerate(src.get('strengths') or [], 1):
            chk(f'{lbl} 強み#{i}', s, 'strength')
        for i, s in enumerate(src.get('priority_issues') or [], 1):
            chk(f'{lbl} 課題#{i}', s, 'issue_top')
    for i, s in enumerate(v.get('top_issues') or [], 1):
        chk(f'C-3 top_issues#{i}', s, 'issue_top')
    chk('C-1 結論（conclusion）', d.get('conclusion'), 'conclusion')
    for i, pr in enumerate(p.get('proposals') or [], 1):
        if not isinstance(pr, dict):
            continue
        chk(f'C-2 提案#{i} title', pr.get('title'), 'p_title')
        chk(f'C-2 提案#{i} target_area', pr.get('target_area'), 'target_area')
        chk(f'C-2 提案#{i} issue', pr.get('issue'), 'p_issue')
        chk(f'C-2 提案#{i} before', pr.get('before'), 'before_after')
        chk(f'C-2 提案#{i} after', pr.get('after'), 'before_after')
    chk('C-2 summary（POINT）', p.get('summary'), 'p_summary')
    chk('C-3 purpose', v.get('purpose'), 'purpose')
    chk('C-3 target', v.get('target'), 'purpose')
    chk('C-3 diagnosis_date', v.get('diagnosis_date'), 'diag_date')
    chk('C-3 総評（summary）', v.get('summary'), 'v_summary')
    for i, sec in enumerate(v.get('sections') or [], 1):
        if not isinstance(sec, dict):
            continue
        a = chk(f'C-3 構造マップ#{i} label', sec.get('label'), 'sec_label')
        b = chk(f'C-3 構造マップ#{i} desc', sec.get('desc'), 'sec_desc')
        if a is not None and b is not None and a + b > 16:
            ng.append(f'C-3 構造マップ#{i}：見出し＋説明の表示換算 {a + b:g}（上限 16）→ 言い換える')
    for i, st in enumerate((v.get('flow_steps') or [])[:6], 1):
        if not isinstance(st, dict):
            continue
        chk(f'C-3 フロー#{i} label', st.get('label'), 'flow_label')
        chk(f'C-3 フロー#{i} note', st.get('note'), 'flow_note')
        comb = ((st.get('label') or '') + ' ' + (st.get('note') or '')).strip()
        if len(comb) > 25:
            ng.append(f'C-3 フロー#{i}：label＋note の len {len(comb)} 超過（上限 25）→ 言い換える')
    chk('C-3 flow_summary', v.get('flow_summary'), 'flow_summary')
    for i, h in enumerate((v.get('highlights') or [])[:3], 1):
        if not isinstance(h, dict):
            continue
        chk(f'C-3 ハイライト#{i} title', h.get('title'), 'hl_title')
        chk(f'C-3 ハイライト#{i} target_area', h.get('target_area'), 'target_area')
        chk(f'C-3 ハイライト#{i} before', h.get('before'), 'before_after')
        chk(f'C-3 ハイライト#{i} after', h.get('after'), 'before_after')
    chk('C-3 改善方向（direction）', v.get('direction'), 'direction')
    if v.get('top_issue') is not None:
        chk('C-3 top_issue', v.get('top_issue'), 'top_issue')

    # ---- §5 record ----
    r = data.get('record')
    if not isinstance(r, dict):
        ng.append('record が無い・型が違う → input.json のいちばん上に record を置く')
    else:
        so, tu, io = r.get('service_name_official'), r.get('target_url', 'MISSING'), r.get('input_origin')
        if not (isinstance(so, str) and so.strip()):
            ng.append('record.service_name_official が無い・空 → 正式名称を入れる')
        if tu == 'MISSING' or not (tu is None or isinstance(tu, str)):
            ng.append('record.target_url が無い・型が違う（文字列か null）')
        if io not in ('url', 'screenshot'):
            ng.append(f'record.input_origin が「{io}」→ url か screenshot のどちらか')
        it = str(d.get('input_type') or '')
        if io == 'url':
            if not (isinstance(tu, str) and re.match(r'^https?://\S+$', tu)):
                ng.append('input_origin が url なのに target_url が空・http(s):// で始まらない')
        if io == 'screenshot' and 'スクショ' not in it:
            ng.append(f'input_origin が screenshot なのに input_type「{it}」に「スクショ」が無い')
        if ('URL' in it) != (io == 'url'):
            ng.append(f'input_type「{it}」と input_origin「{io}」が食い違う（URL を含む ⇔ url）')
        if isinstance(so, str) and so.strip() and names[0] and so != names[0]:
            notes.append(f'表示名「{names[0]}」（正式名称：{so}）')
    return ng, notes


# ---- PPTX の読み出し ----
NSA = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def _emu_px(v):
    return int(v) / EMU_PX


def read_frames(path):
    """スライドごとに、文字枠（表のセルを含む）を図形の順に返す"""
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for sl in prs.slides:
        frames = []
        for sh in sl.shapes:
            box = (sh.left / EMU_PX, sh.top / EMU_PX, sh.width / EMU_PX, sh.height / EMU_PX) \
                if sh.width is not None else None
            if sh.has_text_frame:
                frames.append(dict(kind='text', box=box, body=sh.text_frame._txBody,
                                   text=sh.text_frame.text))
            if getattr(sh, 'has_table', False) and sh.has_table:
                tbl = sh.table
                x = box[0]
                colw = [c.width / EMU_PX for c in tbl.columns]
                y = box[1]
                for ri, row in enumerate(tbl.rows):
                    rh = row.height / EMU_PX
                    cx = x
                    for ci, cell in enumerate(row.cells):
                        tc = cell._tc
                        pr = tc.find(NSA + 'tcPr')
                        mar = {k: (int(pr.get(k)) / EMU_PX if pr is not None and pr.get(k) else d)
                               for k, d in (('marL', 91440 / EMU_PX), ('marR', 91440 / EMU_PX),
                                            ('marT', 45720 / EMU_PX), ('marB', 45720 / EMU_PX))}
                        frames.append(dict(kind='cell', box=(cx, y, colw[ci], rh), body=cell.text_frame._txBody,
                                           text=cell.text_frame.text, row=ri, col=ci, mar=mar,
                                           table_box=box))
                        cx += colw[ci]
                    y += rh
        slides.append(frames)
    return slides


def paragraphs(body):
    """段落ごとに runs[(text, size_pt or None, bold)]・行間・前後間隔・改行の有無"""
    out = []
    for p in body.findall(NSA + 'p'):
        runs, has_br = [], False
        for el in p:
            if el.tag == NSA + 'r':
                rpr = el.find(NSA + 'rPr')
                sz = rpr.get('sz') if rpr is not None else None
                t = el.find(NSA + 't')
                runs.append((t.text or '' if t is not None else '', int(sz) / 100 if sz else None,
                             rpr is not None and rpr.get('b') == '1'))
            elif el.tag == NSA + 'br':
                has_br = True
        ppr = p.find(NSA + 'pPr')
        ln = sa = sb = None
        if ppr is not None:
            x = ppr.find(NSA + 'lnSpc')
            if x is not None and len(x):
                c = x[0]
                ln = ('pct', int(c.get('val')) / 100000) if c.tag.endswith('spcPct') else ('pts', int(c.get('val')) / 100)
            for tag, key in (('spcAft', 'sa'), ('spcBef', 'sb')):
                x = ppr.find(NSA + tag)
                if x is not None and len(x):
                    c = x[0]
                    val = ('pct', int(c.get('val')) / 100000) if c.tag.endswith('spcPct') else ('pts', int(c.get('val')) / 100)
                    if key == 'sa':
                        sa = val
                    else:
                        sb = val
        endp = p.find(NSA + 'endParaRPr')
        end_sz = int(endp.get('sz')) / 100 if endp is not None and endp.get('sz') else None
        out.append(dict(runs=runs, ln=ln, sa=sa, sb=sb, br=has_br, end_sz=end_sz,
                        text=''.join(r[0] for r in runs)))
    return out


def insets(fr):
    if fr['kind'] == 'cell':
        m = fr['mar']; return m['marL'], m['marR'], m['marT'], m['marB']
    bp = fr['body'].find(NSA + 'bodyPr')
    g = lambda k, d: (int(bp.get(k)) if bp is not None and bp.get(k) is not None else d) / EMU_PX
    return g('lIns', 91440), g('rIns', 91440), g('tIns', 45720), g('bIns', 45720)


def para_width(par):
    """段落の幅（px）。(幅, 定義外の文字, 大きさ不明, 最大の大きさ pt, 大きさ混在)"""
    w, undefined, nosize, sizes = 0.0, [], False, set()
    for t, sz, b in par['runs']:
        if not t:
            continue
        if sz is None:
            nosize = True; continue
        sizes.add(sz)
        for ch in t:
            cw = pptx_char_w(ch)
            if cw is None:
                undefined.append(ch)
            else:
                w += sz * 4 / 3 * cw
    mx = max(sizes) if sizes else par.get('end_sz')
    return w, undefined, nosize, mx, len(sizes) > 1


def _spacing_px(val, size_pt):
    if not val:
        return 0.0
    kind, x = val
    return x * 4 / 3 if kind == 'pts' else size_pt * 4 / 3 * x


def long_ascii_word(text, cap_chars):
    for m in re.finditer(r'[!-~]+', text):
        if len(m.group(0)) * 0.5 >= cap_chars / 3:
            return m.group(0)
    return None


# ---- §3・§4 表示の簡易推定 ----
FILE_SLIDES = {'C-1': (2,), 'C-2': (1, 2), 'C-3': (3,)}

# 1行必須の欄（入力の欄）の位置。(ファイル, スライド番号 or None=全, left, top, width, 欄名, 種類)
# 種類：'line'＝1行必須／'multi'＝複数行（高さで点検）／'card'＝見出し＋値の2段落（値は1行・高さも）
#       'flow'＝見出し＋説明の2段落（各1行・合計の高さ）／'map'＝同じ段落内の見出し＋説明（1行）
def frame_role(key, sno, fr, nsl):
    x, y, w, h = [round(v) for v in fr['box']]
    t = fr['text']
    near = lambda a, b: abs(a - b) <= TOL
    if fr['kind'] == 'cell':
        if key == 'C-1' and sno == 2:
            if fr['row'] == 0:
                return ('fixed', '表の見出し', 'line')
            return (('input', f'所見#{fr["row"]}', 'multi_cell') if fr['col'] == 3
                    else ('fixed' if fr['col'] != 2 else 'fixed', f'表 {fr["row"]}行{fr["col"]+1}列', 'line'))
        return None
    if near(x, 40) and near(y, 682) and near(w, 400):
        return ('footer', 'フッター', 'line')
    if near(x, 1140) and near(y, 685) and near(w, 120):
        return ('fixed', 'ページ番号', 'line')
    if near(y, 20) and near(x, 60):
        return ('fixed', '題名', 'line')
    if near(y, 58) and near(x, 60):
        return ('fixed', '副題', 'line')
    if key == 'C-1':
        if near(y, 100) and h >= 60:
            lab = {58: '診断対象', 361: '入力種別', 664: '総合スコア', 967: 'ランク'}
            for k, v in lab.items():
                if near(x, k):
                    return ('input' if v in ('診断対象', '入力種別') else 'fixed', v, 'card')
        if sno == 1:
            if near(y, 190):
                return ('fixed', '欄の見出し', 'line')
            if near(y, 244) and near(x, 60):
                return ('input', '強み（C-1）', 'multi')
            if near(y, 244) and near(x, 668):
                return ('input', '課題（C-1）', 'multi_line')
            if near(y, 594) and near(x, 40):
                return ('fixed', '「結論」', 'line')
            if near(y, 594) and near(x, 168):
                return ('input', '結論', 'multi2')
    if key == 'C-2':
        if near(y, 30) and near(x, 880) or (near(y, 33) and x in range(968, 1133)):
            return ('fixed', '優先度の凡例', 'line')
        if near(x, 58) and near(w, 34):
            return ('fixed', '提案番号', 'line')
        if near(x, 102) and near(w, 770):
            return ('input', '提案 title', 'line')
        if near(x, 980) and near(w, 70):
            return ('fixed', '優先度', 'line')
        if near(x, 1060) and near(w, 170):
            return ('fixed', '工数', 'line')
        if near(x, 56) and near(w, 568):
            return ('input', '改善箇所', 'line')
        if near(x, 640) and near(w, 584):
            return ('input', '課題', 'line')
        if near(x, 64) and near(w, 552):
            return ('input', '現状', 'line')
        if near(x, 648) and near(w, 568):
            return ('input', '改善後', 'line')
        if near(y, 594) and near(x, 40):
            return ('fixed', '「POINT」', 'line')
        if near(y, 594) and near(x, 170):
            return ('input', 'POINT', 'multi2')
    if key == 'C-3':
        if near(y, 98) and h >= 45:
            lab = {58: '目的', 361: '対象', 664: '診断日', 967: '総合スコア'}
            for k, v in lab.items():
                if near(x, k):
                    return ('input' if v != '総合スコア' else 'fixed', v, 'card')
        if near(y, 594) and near(x, 40):
            return ('fixed', '「改善方向」', 'line')
        if near(y, 594) and near(x, 190):
            return ('input', '改善方向', 'multi2')
        if sno == 1:
            if near(y, 158) or (near(y, 164) and w <= 60) or near(y, 238) or near(y, 444) and near(x, 56):
                return ('fixed', '欄の見出し', 'line')
            if near(x, 48) and near(w, 22):
                return ('fixed', '構造マップ番号', 'line')
            if near(x, 78) and near(w, 332):
                return ('input', '構造マップ', 'map')
            if near(x, 446) and near(y, 186):
                return ('input', '総評', 'line')
            if near(x, 446) and near(y, 272):
                return ('input', '最重要課題', 'multi_line')
            if near(x, 920) and near(y, 194):
                return ('input', '強み（C-3 1枚目）', 'multi')
            if near(y, 444) and near(x, 560):
                return ('input', '行動の流れのまとめ', 'line')
            if near(y, 480) and near(w, 44):
                return ('fixed', 'フローの記号', 'line')
            if near(y, 488) and near(w, 12) and t == '→':
                return ('exception', 'フローの矢印', 'arrow')
            if near(y, 528) and near(w, 180):
                return ('input', 'フロー', 'flow')
        if sno == 2:
            if near(y, 158):
                return ('fixed', '欄の見出し', 'line')
            if near(x, 56) and near(w, 180):
                return ('fixed', '10項目の名前', 'line')
            if near(x, 408) and near(w, 62):
                return ('fixed', 'スコア', 'line')
            if near(y, 388):
                return ('fixed', '欄の見出し', 'line')
            if near(x, 656) and near(y, 194):
                return ('input', '強み（C-3 2枚目）', 'multi')
            if near(x, 656) and near(y, 424):
                return ('input', '課題（C-3 2枚目）', 'multi_line')
        if sno == 3:
            if near(x, 60) and near(w, 32):
                return ('fixed', 'ハイライト番号', 'line')
            if near(x, 102) and near(w, 660):
                return ('input', 'ハイライト title', 'line')
            if near(x, 780) and near(w, 70):
                return ('fixed', '優先度', 'line')
            if near(x, 860) and near(w, 90):
                return ('fixed', '工数', 'line')
            if near(x, 960) and near(w, 260):
                return ('input', '箇所', 'line')
            if near(w, 90) and (near(x, 60) or near(x, 644)) and t in ('BEFORE', 'AFTER', '次の一手'):
                return ('fixed', t, 'line')
            if near(x, 158) and near(w, 462):
                return ('input', 'BEFORE', 'line')
            if near(x, 742) and near(w, 462):
                return ('input', 'AFTER', 'line')
    return None


# ---- 判定118：固定カードの高さの限定例外（3種類・6出現）----
# 採用 builder 17.2.17（47629772eaff）の _scorecard_info_cards／_visual_meta_header の構造と一致するときだけ、
# 高さ推定の超過を「要実機確認」へ移す。幅・値・改行・大きさ・枠・段落構造は免除しない
BUILDER_MD5_12 = '47629772eaff'
FIXED_CARD = {
    # (ファイル, 欄名): (枠 x,y,w,h, 見出し, 2段落目の区分[(大きさ,太字,文字の型)], 行間[段落1,段落2])
    ('C-1', '総合スコア'): ((664, 100, 265, 66), '総合スコア',
                         [(32.0, True, r'\d{1,2}'), (14.0, False, r'  / 50')], [1.0, 1.0]),
    ('C-1', 'ランク'): ((967, 100, 265, 66), 'ランク',
                      [(32.0, True, r'[SABCD]'), (14.0, False, r'  (卓越|優良|標準的|要改善|致命的)')], [1.0, 1.0]),
    ('C-3', '総合スコア'): ((967, 98, 265, 50), '総合スコア',
                         [(26.0, True, r'\d{1,2}'), (14.0, False, r'  / 50    '), (26.0, True, r'[SABCD]')], [1.0, 1.0]),
}


def fixed_card_match(key, sno, field, fr, data=None):
    """判定118 の適用条件を満たすか。(True/False, 理由)"""
    spec = FIXED_CARD.get((key, field))
    if spec is None:
        return False, '対象外'
    allowed = {'C-1': (1, 2), 'C-3': (1, 2)}[key]
    if sno not in allowed:
        return False, 'スライドが対象外'
    box, label, segs, lns = spec
    if not all(abs(round(a) - b) <= TOL for a, b in zip(fr['box'], box)):
        return False, '枠の位置・寸法が builder と違う'
    pars = paragraphs(fr['body'])
    if len(pars) != 2:
        return False, f'段落数が {len(pars)}（2であるべき）'
    p1, p2 = pars
    if [(r[0], r[1], r[2]) for r in p1['runs']] != [(label, 14.0, True)]:
        return False, '見出しの文字・大きさ・太字が builder と違う'
    if len(p2['runs']) != len(segs):
        return False, '値の文字区分の数が builder と違う'
    for (t, sz, b), (esz, eb, rx) in zip(p2['runs'], segs):
        if sz != esz or bool(b) != eb or not re.fullmatch(rx, t):
            return False, f'値の区分「{t}」の大きさ・太字・形が builder と違う'
    for par, ln in zip(pars, lns):
        if par['ln'] != ('pct', ln) or par['br'] or par['sa'] or par['sb']:
            return False, '行間・段落の間隔・改行が builder と違う'
    if data:
        src = (data.get('diagnosis') if key == 'C-1' else data.get('visual_data')) or {}
        if field == '総合スコア' and p2['runs'][0][0] != str(src.get('total_score')):
            return False, '総合点が入力と一致しない'
        if key == 'C-1' and field == 'ランク':
            if p2['runs'][0][0] != str(src.get('rank')) or p2['runs'][1][0] != '  ' + str(src.get('rank_label')):
                return False, 'ランク・説明が入力と一致しない'
        if key == 'C-3' and p2['runs'][2][0] != str(src.get('rank')):
            return False, 'ランクが入力と一致しない'
    return True, ''


def _line_h(par, size_pt):
    ln = par['ln']
    if ln is None:
        return size_pt * 4 / 3 * 1.2, True
    kind, x = ln
    return (size_pt * 4 / 3 * 1.2 * x if kind == 'pct' else x * 4 / 3), False


def estimate(fr, kind):
    """1枠の推定。戻り値 dict(status, why, detail)"""
    l, r, t, b = insets(fr)
    x, y, w, h = fr['box']
    if not w or not h:
        return dict(status='未検査', why='枠の幅・高さが取れない')
    cw = w - l - r
    ch = h - t - b
    pars = paragraphs(fr['body'])
    total, total_aux, assumed, notes = 0.0, 0.0, False, []
    width_ng = []
    for i, par in enumerate(pars):
        pw, undefined, nosize, mx, mixed = para_width(par)
        if nosize:
            return dict(status='未検査', why=f'段落{i+1}の文字の大きさが分からない')
        if undefined:
            return dict(status='未検査', why=f'定義されていない文字種：{"".join(sorted(set(undefined)))}')
        if mx is None:
            mx = 14.0
        if mixed:
            notes.append('大きさ混在')
        lh, asm = _line_h(par, mx)
        assumed |= asm
        one_line = kind in ('line', 'map', 'multi_line', 'multi_cell_line') or \
            (kind == 'card' and i == 1) or kind == 'flow'
        if par['br']:
            if one_line:
                width_ng.append(f'段落{i+1}に明示の改行'); 
            notes.append('明示改行あり')
        lines = max(1, -(-pw // cw)) if pw > 0 else 1
        aux = max(1, -(-pw // max(1.0, cw - mx * 4 / 3))) if pw > 0 else 1
        if one_line and pw > cw + 1e-6:
            width_ng.append(f'段落{i+1}の幅 {pw:.1f}px > 枠 {cw:.1f}px「{par["text"]}」')
        lw = long_ascii_word(par['text'], cw / (mx * 4 / 3))
        if lw and not one_line:
            notes.append(f'長い ASCII 単語「{lw}」')
        sp = _spacing_px(par['sb'], mx) + (_spacing_px(par['sa'], mx) if i < len(pars) - 1 else 0)
        total += lines * lh + sp
        total_aux += aux * lh + sp
    res = dict(text=fr['text'], height=round(total, 2), cap=round(ch, 2), notes=notes)
    if width_ng:
        res.update(status='NG', why='幅 NG：' + '／'.join(width_ng))
        if total > ch + 1e-6:
            res['why'] += f'／同じ枠の高さ超過 {total:.1f}px > {ch:.1f}px'
            res['height_ng'] = True
        return res
    # §3-2 の対象＝複数行を許す欄と、複数の段落がある枠（1段落・1行の枠は幅で点検）
    height_scope = kind in ('multi', 'multi_line', 'card', 'flow') or len(pars) >= 2
    if height_scope and total > ch + 1e-6:
        res.update(status='NG', why=f'高さ超過 {total:.1f}px > 枠 {ch:.1f}px', height_ng=True)
        return res
    if not height_scope:
        total_aux = 0.0      # 1段落・1行の枠は幅で点検する（高さ・補助の判定の対象外）
    reasons = []
    if total_aux > ch + 1e-6:
        reasons.append(f'折り返しの余裕不足（補助の判定 {total_aux:.1f}px > {ch:.1f}px）')
    if assumed and height_scope:
        notes.append('行間：仮定1.0')
        if (ch - total) / ch < 0.10:
            reasons.append('行間を仮定し高さの余裕10%未満')
    if any(n.startswith('長い ASCII') for n in notes):
        reasons.append('長い ASCII 単語')
    if reasons:
        res.update(status='要実機確認', why='／'.join(reasons))
    else:
        res.update(status='基準内', why='')
    return res


def check_display(found, data=None):
    """found={'C-1':path,...}。欄ごとの結果リストと、フッター・枚数の NG"""
    rows, ng = [], []
    for key in ('C-1', 'C-2', 'C-3'):
        if key not in found:
            continue
        fn = os.path.basename(found[key])
        try:
            slides = read_frames(found[key])
        except Exception as e:
            rows.append(dict(file=fn, slide=0, field='（ファイル）', status='未検査', why=f'PPTX が開けない：{e}'))
            continue
        if len(slides) not in FILE_SLIDES[key]:
            rows.append(dict(file=fn, slide=0, field='（枚数）', status='未検査',
                             why=f'{len(slides)}枚（想定 {"・".join(map(str, FILE_SLIDES[key]))}枚）'))
            continue
        for sno, frames in enumerate(slides, 1):
            footer_n = 0
            for fr in frames:
                if not fr['text'].strip():
                    continue
                if 'by GPTs' in fr['text']:
                    ng.append(f'{fn} {sno}枚目：「by GPTs」を含む枠がある「{fr["text"]}」')
                role = frame_role(key, sno, fr, len(slides))
                base = dict(file=fn, slide=sno, box=[round(v) for v in fr['box']], text=fr['text'])
                if role is None:
                    rows.append(dict(base, field='（未知の枠）', status='未検査',
                                     why='入力の欄・固定文言・例外のどれにも結び付かない'))
                    continue
                cls, field, kind = role
                if cls == 'footer':
                    footer_n += 1
                    st = '基準内' if fr['text'] == AUTHOR else 'NG'
                    rows.append(dict(base, field=field, status=st, kind='footer',
                                     why='' if st == '基準内' else f'表記の不一致「{fr["text"]}」（正：{AUTHOR}）'))
                    continue
                if cls == 'exception':
                    pw = para_width(paragraphs(fr['body'])[0])[0]
                    rows.append(dict(base, field=field, status='参考', kind='arrow',
                                     why=f'例外表 #1（幅 {pw:.1f}px／枠 {fr["box"][2]:.0f}px）・実機確認'))
                    continue
                kk = {'multi_cell': 'multi', 'multi_line': 'multi_line', 'multi2': 'multi'}.get(kind, kind)
                if kind == 'multi_line':
                    kk = 'multi_line'
                r = estimate(fr, kk)
                if (key, field) in FIXED_CARD:
                    okc, whyc = fixed_card_match(key, sno, field, fr, data)
                    if not okc:
                        r = dict(r, status='NG' if r.get('status') != '未検査' else '未検査',
                                 why=(r.get('why') + '／' if r.get('why') else '') + f'固定カードの構造が builder と違う：{whyc}')
                    elif r.get('status') == 'NG' and r.get('height_ng') and not str(r.get('why', '')).startswith('幅 NG'):
                        r = dict(r, status='要実機確認', exception='固定カード',
                                 why=f'固定カード／高さ推定{r["height"]}px・枠{r["cap"]:g}px／高さ判定は限定例外（判定118）'
                                     f'／Windows PowerPoint未確認')
                    elif r.get('status') == '基準内':
                        r = dict(r, why='固定カード（構造一致）')
                rows.append(dict(base, field=field, kind=kind, cls=cls, **{k: v for k, v in r.items() if k != 'text'}))
            if footer_n != 1:
                ng.append(f'{fn} {sno}枚目：フッターが{footer_n}件（1件であるべき）')
    return rows, ng


def summarize(rows):
    cnt = {k: [r for r in rows if r['status'] == k] for k in ('基準内', '要実機確認', 'NG', '未検査', '参考')}
    return cnt


def print_display(rows, extra_ng):
    cnt = summarize(rows)
    print(f'表示の簡易推定：基準内（{len(cnt["基準内"])}欄）／要実機確認（{len(cnt["要実機確認"])}欄）'
          f'／NG（{len(cnt["NG"]) + len(extra_ng)}件）／未検査（{len(cnt["未検査"])}欄）／参考（{len(cnt["参考"])}件）')
    for k in ('NG', '未検査', '要実機確認', '参考'):
        for r in cnt[k]:
            print(f'  [{k}] {r["file"]} {r["slide"]}枚目 {r["field"]}：{r.get("why", "")}')
    for m in extra_ng:
        print(f'  [NG] {m}')
    return len(cnt['NG']) + len(extra_ng), len(cnt['未検査'])


def check_pptx_vs_input(data, found):
    """§3-4：入力の文字列が切り詰めなしで PPTX にあるか（末尾「…」の検出を含む）"""
    ng = []
    d, p, v = data.get('diagnosis') or {}, data.get('proposals') or {}, data.get('visual_data') or {}
    def texts(key):
        if key not in found:
            return []
        return [fr['text'] for sl in read_frames(found[key]) for fr in sl]
    t1, t2, t3 = texts('C-1'), texts('C-2'), texts('C-3')
    for key, tx in (('C-1', t1), ('C-2', t2), ('C-3', t3)):
        for t in tx:
            if t.rstrip().endswith('…'):
                ng.append(f'{key}：末尾が「…」の枠がある（builder が切り詰めた）「{t}」')
    def has(tx, s):
        return any(s in t for t in tx)
    want = []
    want += [('C-1', t1, x) for x in [d.get('service_name'), d.get('input_type'), d.get('conclusion')]]
    want += [('C-1', t1, '● ' + s) for s in d.get('strengths') or []]
    for pr in p.get('proposals') or []:
        want += [('C-2', t2, pr.get('title')), ('C-2', t2, '改善箇所　' + str(pr.get('target_area'))),
                 ('C-2', t2, '課題　' + str(pr.get('issue'))), ('C-2', t2, '現状　' + str(pr.get('before'))),
                 ('C-2', t2, '改善後　' + str(pr.get('after')))]
    want += [('C-3', t3, x) for x in [v.get('purpose'), v.get('target'), v.get('diagnosis_date')]]
    want += [('C-3', t3, '● ' + s) for s in v.get('strengths') or []]
    for sec in v.get('sections') or []:
        want.append(('C-3', t3, f'{sec.get("label")}  {sec.get("desc")}'))
    for st in (v.get('flow_steps') or [])[:6]:
        want.append(('C-3', t3, f'{st.get("label")}\n{st.get("note")}'))
    for h in (v.get('highlights') or [])[:3]:
        want += [('C-3', t3, h.get('title')), ('C-3', t3, '箇所：' + str(h.get('target_area'))),
                 ('C-3', t3, h.get('before')), ('C-3', t3, h.get('after'))]
    want.append(('C-3', t3, (v.get('direction') or '').replace('\n', ' ')))
    for key, tx, s in want:
        if s and tx and not has(tx, s):
            ng.append(f'{key}：入力の文字「{s}」が PPTX にそのまま無い（切り詰め・作り直し漏れ）')
    return ng


def find_files(out_dir):
    files = sorted(os.path.basename(x) for x in glob.glob(os.path.join(out_dir, '*.pptx')))
    found = {}
    for key, rx in NAME_RE.items():
        hit = [f for f in files if rx.match(f)]
        if len(hit) == 1:
            found[key] = os.path.join(out_dir, hit[0])
    return found


def main(argv):
    report = None
    if '--report' in argv:
        i = argv.index('--report'); report = argv[i + 1]; argv = argv[:i] + argv[i + 2:]
    if len(argv) == 2 and argv[0] == '--pptx-only':
        found = find_files(argv[1])
        rows, extra = check_display(found)
        print('点検対象：', ', '.join(os.path.basename(v) for v in found.values()) or '（なし）')
        print('内容・構成点検：未実施（入力なし・--pptx-only）')
        n_ng, n_un = print_display(rows, extra)
        print('Windows PowerPoint 目視：未実施')
        print('※「基準内」は実際の描画の合格を意味しない。最終の判定は Windows の PowerPoint での目視。')
        if report:
            json.dump(dict(rows=rows, extra_ng=extra), open(report, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        return 3 if (n_ng == 0 and n_un == 0) else 1
    if len(argv) != 2:
        print(__doc__); return 2
    data = json.load(open(argv[0], encoding='utf-8'))
    files, ng = check(data, argv[1])                  # r5 の点検（そのまま）
    ng_in, notes = check_input_r6(data)               # r6：入力の A・B・C・文字種・record
    found = find_files(argv[1])
    try:
        ng_px = check_pptx_vs_input(data, found)      # r6：切り詰め・入力との一致
    except Exception as e:
        ng_px = [f'PPTX の読み出しに失敗：{e}']
    rows, extra = check_display(found, data)          # r6：表示の簡易推定（入力の結果に関わらず実行）
    content_ng = ng + ng_in + ng_px
    print('点検対象：', ', '.join(files) or '（なし）')
    for n in notes:
        print('  表示名の報告：', n)
    if content_ng:
        print(f'内容・構成点検：🚨 不合格 {len(content_ng)}件')
        for m in content_ng:
            print('  -', m)
    else:
        print('内容・構成点検：⭕ 合格（総評17字・ランクの説明・件数・ファイル名・test_*.pptx 無し・リンク無し・'
              'メイリオ14pt以上・PPTX と入力の一致・Before/After の箇所と優先度・行動の流れのまとめ・'
              '字数 A/B/C・文字種・record）')
    n_ng, n_un = print_display(rows, extra)
    print('Windows PowerPoint 目視：未実施')
    print('※「基準内」は実際の描画の合格を意味しない。最終の判定は Windows の PowerPoint での目視。')
    if report:
        json.dump(dict(content_ng=content_ng, rows=rows, extra_ng=extra), open(report, 'w', encoding='utf-8'),
                  ensure_ascii=False, indent=1)
    ok = (not content_ng) and n_ng == 0 and n_un == 0
    print('⭕ 点検合格（表示は簡易推定・実機の目視は別途）' if ok else '🚨 不合格 → 入力を直して PPTX を作り直す')
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
