# 判定126 の追加試験（r6.2）
import json, copy, os, shutil, subprocess, sys, importlib.util
from pptx import Presentation
from pptx.util import Pt, Emu
from lxml import etree
R6=sys.argv[1] if len(sys.argv)>1 else '/tmp/r6work/r62/ui-diagnosis-director'
CHK=R6+'/scripts/check_output.py'
W='/tmp/r6work/bk2'; shutil.rmtree(W,ignore_errors=True); os.makedirs(W)
BASE=json.load(open(R6+'/assets/sample_input_rankB.json',encoding='utf-8'))
A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
F3='03_Visual_Board_sample_lp.pptx'; F1='01_UI_Scorecard_sample_lp.pptx'
rows=[]
def gen(data,d):
    jp=d+'.json'; json.dump(data,open(jp,'w',encoding='utf-8'),ensure_ascii=False)
    subprocess.run(['python3','/tmp/r6work/harness_r6.py',R6,jp,d],capture_output=True,check=True); return jp
def check(jp,d): c=subprocess.run(['python3',CHK,jp,d],capture_output=True,text=True); return c.returncode,c.stdout
def T(name, data_mut, pptx_mut, fname, want, expect_exit, notwant=None, extra=None):
    d=f'{W}/{len(rows):02d}'; data=copy.deepcopy(BASE); data_mut(data); jp=gen(data,d)
    if pptx_mut:
        p=os.path.join(d,fname); prs=Presentation(p); pptx_mut(prs); prs.save(p)
    rc,out=check(jp,d)
    hit=[l.strip() for l in out.splitlines() if all(w in l for w in want) and not (notwant and notwant in l)]
    ok=rc==expect_exit and bool(hit) and (extra is None or extra(out))
    rows.append(('⭕' if ok else '🚨',name,rc,(hit[0] if hit else '(期待の行なし)')[:170]))
def nums(prs):
    return [s for s in prs.slides[2].shapes if s.has_text_frame and s.text_frame.text.strip() and abs(s.left/9525-60)<=2 and abs(s.width/9525-32)<=2]
def setnum(prs,i,t): nums(prs)[i].text_frame.paragraphs[0].runs[0].text=t
nobad=lambda out: not any('3-3-HL' in l and '番号' in l and ('[NG]' in l or '[未検査]' in l) for l in out.splitlines())
# ---- 126-1：ハイライト番号 ----
T('番号：no 未指定・出現順（正常）', lambda d: None, None, F3, ['点検合格'], 0, extra=nobad)
T('番号：no 明示 1・2・3（正常）', lambda d: [h.update(no=i+1) for i,h in enumerate(d['visual_data']['highlights'])], None, F3, ['点検合格'], 0, extra=nobad)
def nonseq(d):
    for h,n in zip(d['visual_data']['highlights'],(1,3,5)): h['no']=n
T('番号：正当な非連番 no=1・3・5', nonseq, None, F3, ['点検合格'], 0, extra=nobad)
def ref(d):
    hs=d['visual_data']['highlights']; ti=d['visual_data']['top_issues']
    hs[0]['issue_ref']=ti[2]; hs[2]['issue_ref']=ti[0]   # 1件目は論点3、3件目は論点1を指す
T('番号：no 未指定・論点照合で 3・2・1', ref, None, F3, ['点検合格'], 0, extra=nobad)
T('番号：誤番号 1→9（no 未指定）', lambda d: None, lambda P: setnum(P,0,'9'), F3, ['3-3-HL1-番号','固定文言の文字が期待と一致しない'], 1)
T('番号：誤番号 1→9（no=1 明示）', lambda d: d['visual_data']['highlights'][0].update(no=1), lambda P: setnum(P,0,'9'), F3, ['3-3-HL1-番号','一致しない'], 1)
T('番号：非連番 1・3・5 の入力に 1・2・3 の PPTX', nonseq, lambda P: (setnum(P,1,'2'),setnum(P,2,'3')), F3, ['3-3-HL2-番号','一致しない'], 1)
T('番号：論点照合 3 の箇所を 1 に', ref, lambda P: setnum(P,0,'1'), F3, ['3-3-HL1-番号','一致しない'], 1)
# ---- 126-2：フッター ----
def foot(prs,sl=0):
    return [s for s in prs.slides[sl].shapes if s.has_text_frame and s.text_frame.text=='UI診断ディレクター'][0]
def spc(v):
    def f(P):
        for r in foot(P).text_frame.paragraphs[0].runs: r._r.get_or_add_rPr().set('spc',str(v))
    return f
T('フッター：正常', lambda d: None, None, F1, ['点検合格'], 0)
T('フッター：小さな文字間隔（spc=100）で枠内', lambda d: None, spc(100), F1, ['点検合格'], 0)
T('フッター：大きな文字間隔（spc=5000）で幅超過', lambda d: None, spc(5000), F1, ['1-1-フッター','幅 NG'], 1)
T('フッター：文字の大きさ 40pt で幅超過', lambda d: None, lambda P: setattr(foot(P).text_frame.paragraphs[0].runs[0].font,'size',Pt(40)), F1, ['1-1-フッター','幅 NG'], 1)
def lst(P):
    l=foot(P).text_frame._txBody.find(A+'lstStyle'); lv=etree.SubElement(l,A+'lvl1pPr'); dr=etree.SubElement(lv,A+'defRPr'); dr.set('spc','100')
T('フッター：継承された文字間隔（解決不能）', lambda d: None, lst, F1, ['1-1-フッター','未検査','継承'], 1)
T('フッター：文言変更', lambda d: None, lambda P: setattr(foot(P).text_frame.paragraphs[0].runs[0],'text','UI診断ディレクター版'), F1, ['1-1-フッター','表記の不一致'], 1)
def brk(P):
    r=foot(P).text_frame.paragraphs[0]._p.findall(A+'r')[-1]; r.addnext(etree.Element(A+'br'))
T('フッター：明示改行', lambda d: None, brk, F1, ['1-1-フッター','明示の改行'], 1)
def both(P):
    rr=foot(P).text_frame.paragraphs[0].runs[0]; rr.text='UI/UX診断 by GPTs'; rr._r.get_or_add_rPr().set('spc','5000')
T('フッター：文言不一致と幅超過が同時（理由を両方残す）', lambda d: None, both, F1, ['1-1-フッター','表記の不一致','幅 NG'], 1)
def remove_footer(P): s=foot(P,1); s._element.getparent().remove(s._element)
T('フッター：2枚目のフッター削除（数の検査を維持）', lambda d: None, remove_footer, F1, ['1-2-フッター','見つからない'], 1)
def move_footer(P): foot(P).left=Emu(600*9525)
T('フッター：位置の移動', lambda d: None, move_footer, F1, ['想定外の枠'], 1)
for r in rows: print(r[0],r[1],'| exit',r[2],'|',r[3])
print('---',sum(1 for r in rows if r[0]=='⭕'),'/',len(rows))
