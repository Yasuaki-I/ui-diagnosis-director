# 判定122 の崩し試験（r6.1）。正常な B 見本の生成物を /tmp で複写し、PPTX をメモリ上で改変して点検する
import json, copy, os, shutil, subprocess, sys
from pptx import Presentation
from pptx.util import Pt, Emu
from lxml import etree
R6=sys.argv[1] if len(sys.argv)>1 else '/tmp/r6work/r62/ui-diagnosis-director'
CHK=R6+'/scripts/check_output.py'
W='/tmp/r6work/bk'; shutil.rmtree(W,ignore_errors=True); os.makedirs(W)
INP=R6+'/assets/sample_input_rankB.json'
BASEDIR=W+'/base'
subprocess.run(['python3','/tmp/r6work/harness_r6.py',R6,INP,BASEDIR],capture_output=True,check=True)
F={'1':'01_UI_Scorecard_sample_lp.pptx','2':'02_Proposals_sample_lp.pptx','3':'03_Visual_Board_sample_lp.pptx'}
A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
rows=[]
def shape(prs,sl,startswith=None,eq=None):
    for s in prs.slides[sl].shapes:
        if s.has_text_frame and ((eq is not None and s.text_frame.text==eq) or (startswith and s.text_frame.text.startswith(startswith))):
            return s
    raise KeyError((sl,startswith,eq))
def T(name, fileno, mut, want, data_mut=None, expect_exit=1, notwant=None):
    d=f'{W}/{len(rows):02d}'; shutil.copytree(BASEDIR,d)
    p=os.path.join(d,F[fileno]); prs=Presentation(p); mut(prs); prs.save(p)
    jp=INP
    if data_mut:
        data=json.load(open(INP,encoding='utf-8')); data_mut(data); jp=d+'.json'; json.dump(data,open(jp,'w',encoding='utf-8'),ensure_ascii=False)
    c=subprocess.run(['python3',CHK,jp,d],capture_output=True,text=True)
    hit=[l.strip() for l in c.stdout.splitlines() if all(w in l for w in want) and not (notwant and notwant in l)]
    ok=(c.returncode==expect_exit) and bool(hit)
    rows.append(('⭕' if ok else '🚨',name,c.returncode,(hit[0] if hit else '(期待の理由なし)')[:170]))
def dup(prs,sl,s):
    el=copy.deepcopy(s._element); prs.slides[sl].shapes._spTree.append(el)
def remove(prs,sl,s): s._element.getparent().remove(s._element)
def settext(s,t): s.text_frame.paragraphs[0].runs[0].text=t
# ---- 122-1：期待枠の欠落・重複・固定文言の変更・想定外 ----
T('固定見出し「◤ LP構造マップ」を削除','3',lambda P: remove(P,0,shape(P,0,eq='◤ LP構造マップ')),['3-1-構造マップ見出し','見つからない'])
T('固定見出しを改変（◤ 構造マップ）','3',lambda P: settext(shape(P,0,eq='◤ LP構造マップ'),'◤ 構造マップ'),['3-1-構造マップ見出し','固定文言の文字が期待と一致しない'])
T('固定見出しを複製','3',lambda P: dup(P,0,shape(P,0,eq='◤ LP構造マップ')),['3-1-構造マップ見出し','重複'])
T('可変枠（総評）を複製','3',lambda P: dup(P,0,shape(P,0,eq='FV内CTA追加で無料体験へ誘導')),['3-1-総評','重複'])
def swap(P):
    a=shape(P,0,startswith='ファーストビュー  '); b=shape(P,0,startswith='特徴  ')
    ta,tb=a.text_frame.paragraphs[0].runs[0].text,b.text_frame.paragraphs[0].runs[0].text
    a.text_frame.paragraphs[0].runs[0].text=tb; b.text_frame.paragraphs[0].runs[0].text=ta
T('可変枠を入替え（構造マップ1と2の見出し）','3',swap,['3-1-構造マップ1','入力の欄の文字が期待と一致しない'])
def addframe(P):
    tb=P.slides[0].shapes.add_textbox(Emu(600*9525),Emu(300*9525),Emu(200*9525),Emu(30*9525)); tb.text_frame.text='余計な枠'
T('想定外の枠を追加','3',addframe,['想定外の枠'])
T('枠を別の位置へ移動（強み見出し）','1',lambda P: setattr(shape(P,0,eq='◎ 強み（活かすべき点）'),'left',Emu(300*9525)),['1-1-強み見出し','見つからない'])
T('ページ番号を改変（1 / 2→1 / 3）','1',lambda P: settext(shape(P,0,eq='1 / 2'),'1 / 3'),['ページ番号','一致しない'])
T('表のセルの文字を改変（項目名）','1',lambda P: setattr(P.slides[1].shapes[[i for i,s in enumerate(P.slides[1].shapes) if s.has_table][0]].table.cell(1,1),'text','別の項目'),['1-2-表-1-項目','一致しない'])
# ---- 122-2：矢印の例外の全条件 ----
def arrows(P): return [s for s in P.slides[0].shapes if s.has_text_frame and s.text_frame.text=='→']
T('矢印を x9999・高さ1 へ','3',lambda P: (setattr(arrows(P)[0],'left',Emu(9999*9525)),setattr(arrows(P)[0],'height',Emu(1*9525))),['想定外の枠'])
T('矢印の高さだけ1に','3',lambda P: setattr(arrows(P)[0],'height',Emu(1*9525)),['想定外の枠'])
T('矢印の大きさを24ptに','3',lambda P: setattr(arrows(P)[0].text_frame.paragraphs[0].runs[0].font,'size',Pt(24)),['矢印','条件を満たさない','大きさ'])
T('矢印を太字なしに','3',lambda P: setattr(arrows(P)[0].text_frame.paragraphs[0].runs[0].font,'bold',False),['矢印','条件を満たさない'])
T('矢印の文字を「⇒」に','3',lambda P: settext(arrows(P)[0],'⇒'),['矢印','条件を満たさない'])
T('矢印を「→→」に','3',lambda P: settext(arrows(P)[0],'→→'),['矢印','条件を満たさない'])
T('矢印を1つ削除（個数）','3',lambda P: remove(P,0,arrows(P)[2]),['矢印3','見つからない'])
T('矢印を複製（個数）','3',lambda P: dup(P,0,arrows(P)[0]),['矢印1','重複'])
def arrow_spc(P):
    arrows(P)[0].text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set('spc','300')
T('矢印に文字間隔','3',arrow_spc,['矢印','文字間隔'])
# 固定カード：余白・構造を変えても免除が広がらない
def card(P): return shape(P,0,startswith='総合スコア')
T('固定カードの左余白を変更','1',lambda P: setattr(card(P).text_frame,'margin_left',Emu(5*9525)),['総合スコア','内側の余白'])
T('固定カードの上余白を変更','1',lambda P: setattr(card(P).text_frame,'margin_top',Emu(5*9525)),['総合スコア','内側の余白'])
def card3(P):
    s=card(P); s.text_frame.add_paragraph().text='追加'
T('固定カードに段落を追加','1',card3,['総合スコア'])
def card_spc(P): card(P).text_frame.paragraphs[1].runs[0]._r.get_or_add_rPr().set('spc','200')
T('固定カードの数字に文字間隔','1',card_spc,['総合スコア','builder と違う'])
# ---- 122-3：C-2／C-3 の対応 ----
T('入力：ハイライト1の優先度 高→低','3',lambda P: None,['優先度が一致しない'],data_mut=lambda d: d['visual_data']['highlights'][0].update(priority='低'))
T('入力：ハイライト1の工数 小→大','3',lambda P: None,['工数が一致しない'],data_mut=lambda d: d['visual_data']['highlights'][0].update(effort='大'))
T('入力：ハイライト2の after を別文','3',lambda P: None,['改善後（after）が一致しない'],data_mut=lambda d: d['visual_data']['highlights'][1].update(after='別の改善後'))
T('入力：ハイライト3の箇所を対応なしに','3',lambda P: None,['対応する C-2 の提案が無い'],data_mut=lambda d: d['visual_data']['highlights'][2].update(target_area='存在しない箇所'))
def dup_area(d):
    d['proposals']['proposals'][3]['target_area']=d['proposals']['proposals'][0]['target_area']
T('入力：同じ箇所の提案が2つ（複数対応）','3',lambda P: None,['複数ある'],data_mut=dup_area)
T('入力：ハイライト1と2が同じ提案に対応','3',lambda P: None,['同じ提案'],data_mut=lambda d: d['visual_data']['highlights'][1].update(**{k:d['visual_data']['highlights'][0][k] for k in ('target_area','priority','effort','before','after')}))
T('PPTX：C-3 の優先度だけ入替え（高→低）','3',lambda P: settext(shape(P,2,eq='高'),'低'),['優先度'])
T('PPTX：C-2 の改善後の末尾に余計な文字','2',lambda P: setattr(shape(P,0,startswith='改善後\u3000無料体験').text_frame.paragraphs[0].runs[1],'text','無料体験と料金確認の2ボタンをFV内に置く！'),['改善後'])
T('PPTX：C-3 の現状を別欄（AFTER）へ移動','3',lambda P: settext(shape(P,2,eq='見出しと説明文のみで主要ボタンがない'),'無料体験と料金確認の2ボタンをFV内に置く'),['現状'])
T('PPTX：C-3 の箇所の枠を削除','3',lambda P: remove(P,2,shape(P,2,eq='箇所：ファーストビュー')),['箇所'])
# ---- 122-4：文字間隔・明示改行 ----
def spc_on(sh,v):
    for r in sh.text_frame.paragraphs[-1].runs: r._r.get_or_add_rPr().set('spc',str(v))
T('正常な枠に小さな文字間隔（総評 spc=50）→ 収まる','3',lambda P: spc_on(shape(P,0,eq='FV内CTA追加で無料体験へ誘導'),50),['点検合格'],expect_exit=0)
T('文字間隔で幅を超える（総評 spc=1000）','3',lambda P: spc_on(shape(P,0,eq='FV内CTA追加で無料体験へ誘導'),1000),['3-1-総評','幅 NG'])
def lst_spc(P):
    s=shape(P,0,eq='FV内CTA追加で無料体験へ誘導'); lst=s.text_frame._txBody.find(A+'lstStyle')
    lvl=etree.SubElement(lst,A+'lvl1pPr'); dr=etree.SubElement(lvl,A+'defRPr'); dr.set('spc','100')
T('継承された文字間隔（lstStyle）','3',lambda P: lst_spc(P),['3-1-総評','継承された文字間隔'])
def add_br(sh, idx_para=0, times=1, empty=False):
    p=sh.text_frame.paragraphs[idx_para]._p
    r=p.findall(A+'r')[-1]
    for _ in range(times):
        br=etree.Element(A+'br'); r.addnext(br)
        if empty: pass
T('1行必須の欄に改行（総評）','3',lambda P: add_br(shape(P,0,eq='FV内CTA追加で無料体験へ誘導')),['3-1-総評','明示の改行'])
def concl_br(P,times,empty=False):
    s=shape(P,0,startswith='FV内CTAと成果訴求'); add_br(s,0,times)
T('複数行の欄に改行1回（結論）→ 高さは2行で収まる・入力との不一致で NG','1',lambda P: concl_br(P,1),['1-1-結論','入力の欄の文字が期待と一致しない'],notwant='高さ超過')
T('複数行の欄に改行3回（結論・行が増えて超過）','1',lambda P: concl_br(P,3),['1-1-結論','高さ超過'])
def concl_mixed(P):
    s=shape(P,0,startswith='FV内CTAと成果訴求'); p=s.text_frame.paragraphs[0]._p
    r=p.findall(A+'r')[-1]; br=etree.Element(A+'br'); r.addnext(br)
    r2=copy.deepcopy(r); r2.find(A+'t').text='大'; r2.find(A+'rPr').set('sz','3200'); br.addnext(r2)
T('改行の後に大きい文字（結論・32pt）','1',concl_mixed,['1-1-結論','高さ超過'])
for r in rows: print(r[0],r[1],'| exit',r[2],'|',r[3])
print('---',sum(1 for r in rows if r[0]=='⭕'),'/',len(rows))
