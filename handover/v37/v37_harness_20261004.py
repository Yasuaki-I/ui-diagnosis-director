# -*- coding: utf-8 -*-
"""v3.7 検証ハーネス：builder を指定して固定入力セットから .pptx を生成し、図形座標・テキストを TSV へ出力する。
使い方: python3 v37_harness.py <builder.py> <out_prefix>  → <out_prefix>.pptx / <out_prefix>.tsv"""
import importlib.util, sys, json
from pptx import Presentation
sys.path.insert(0, __import__('os').path.dirname(__file__))
from v37_fixtures_20261004 import CASES
b, out = sys.argv[1], sys.argv[2]
spec = importlib.util.spec_from_file_location('bld', b); m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
prs = m.create_presentation(); pal = m.get_theme_palette('LightBlue')
reps = []
for i, (k, p, d) in enumerate(CASES):
    _, r = m.add_diagram_slide(prs, p, pal, json.loads(json.dumps(d)), page_num=i+1, total=len(CASES))
    reps.append((k, r))
prs.save(out + '.pptx')
E = 9525
with open(out + '.tsv', 'w') as f:
    for si, s in enumerate(Presentation(out + '.pptx').slides):
        for j, sh in enumerate(s.shapes):
            t = sh.text_frame.text.replace('\n', '\\n') if sh.has_text_frame else ''
            f.write('%s\t%d\t%s\t%s\t%s\t%s\t%s\t%s\n' % (CASES[si][0], j, sh.shape_type, round(sh.left/E,2), round(sh.top/E,2), round(sh.width/E,2), round(sh.height/E,2), t))
with open(out + '.notes.json', 'w') as f: json.dump(reps, f, ensure_ascii=False, indent=1, default=str)
print('cases', len(CASES), '→', out + '.pptx')
