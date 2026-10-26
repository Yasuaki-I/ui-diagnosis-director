import json, copy, os, shutil, subprocess
from pptx import Presentation
R6='/tmp/r6work/r61/ui-diagnosis-director'; W='/tmp/r6work/bt2'; shutil.rmtree(W,ignore_errors=True); os.makedirs(W)
BASE=json.load(open(R6+'/assets/sample_input_rankB.json',encoding='utf-8'))
def run(data,tag,mutate_pptx=None):
    d=f'{W}/{tag}'; jp=d+'.json'; json.dump(data,open(jp,'w',encoding='utf-8'),ensure_ascii=False)
    g=subprocess.run(['python3','/tmp/r6work/harness_r6.py',R6,jp,d],capture_output=True,text=True)
    if g.returncode: return 'builder停止',g.stderr[-200:]
    if mutate_pptx: mutate_pptx(d)
    c=subprocess.run(['python3',R6+'/scripts/check_output.py',jp,d],capture_output=True,text=True)
    return c.returncode,c.stdout
# 異体字（理由表示の確認）
import importlib.util
spec=importlib.util.spec_from_file_location('co',R6+'/scripts/check_output.py'); co=importlib.util.module_from_spec(spec); spec.loader.exec_module(co)
d=copy.deepcopy(BASE); d['visual_data']['summary']='葛\U000E0100'; print('異体字U+E0100 理由:',co.check_input_r6(d)[0])
# 複数行の欄を枠の行数より1行多く（PPTX の結論に3行分の文字を入れる）
def three_lines(dirp):
    p=dirp+'/01_UI_Scorecard_sample_lp.pptx'; prs=Presentation(p)
    sh=[s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.startswith('FV内CTAと')][0]
    sh.text_frame.paragraphs[0].runs[0].text='結'*170; prs.save(p)
rc,out=run(copy.deepcopy(BASE),'height_over',three_lines)
print('複数行1行超過:', 'exit',rc, [l.strip() for l in out.splitlines() if '結論' in l and '[' in l][:1])
# ランク S・A・C・D のカード内容（判定118 の境界）
cats=[x['category'] for x in BASE['diagnosis']['scores']]
for rank,label,score in (('S','卓越',5),('A','優良',4),('C','要改善',2),('D','致命的',1)):
    d=copy.deepcopy(BASE); tot=score*10
    for k in ('diagnosis','visual_data'):
        d[k].update(rank=rank,rank_label=label,total_score=tot)
        for x in d[k]['scores']: x['score']=score
    rc,out=run(d,'card_'+rank)
    cards=[l.strip()[:90] for l in out.splitlines() if '固定カード' in l]
    others=[l.strip()[:120] for l in out.splitlines() if l.strip().startswith('[NG]') or l.strip().startswith('[未検査]')]
    print(f'カード ランク{rank}（{tot}点・{label}）: exit',rc,'固定カード要実機確認',len(cards),'NG/未検査',others[:3])
# 総合点1桁（例：9点＝ランクD）
d=copy.deepcopy(BASE)
for k in ('diagnosis','visual_data'):
    d[k].update(rank='D',rank_label='致命的',total_score=9)
rc,out=run(d,'card_1digit'); print('カード 1桁（9点）: exit',rc,sum('固定カード' in l for l in out.splitlines()),[l.strip()[:100] for l in out.splitlines() if '[NG]' in l][:2])
